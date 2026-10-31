"""Module A – CV Parser: đọc văn bản, chuẩn hoá, tách mục, số năm kinh nghiệm, học vấn.

Việc trích kỹ năng nằm ở `extract.py`; `parse_cv` ghép các bước lại.
"""
from __future__ import annotations

import io
import re
import unicodedata
from datetime import date
from typing import Optional

from .types import CvError

MAX_CV_SIZE = 5 * 1024 * 1024  # trùng MAX_CV_SIZE ở src/utils/cvParser.ts
MIN_TEXT_CHARS = 30

# ---------------------------------------------------------------------------
# 1. Đọc văn bản
# ---------------------------------------------------------------------------


def extract_text(data: bytes, filename: str) -> str:
    """Đọc văn bản từ PDF (có lớp chữ), DOCX hoặc TXT. CV dạng ảnh scan báo lỗi `no_text`."""
    if len(data) > MAX_CV_SIZE:
        raise CvError("too_large", "Tệp lớn hơn 5 MB.")
    name = filename.lower()
    if name.endswith(".pdf"):
        text = _read_pdf(data)
    elif name.endswith(".docx"):
        text = _read_docx(data)
    elif name.endswith(".txt"):
        text = data.decode("utf-8", errors="ignore")
    else:
        raise CvError("unsupported_type", "Chỉ hỗ trợ tệp PDF, DOCX hoặc TXT.")
    if len(re.sub(r"\s", "", text)) < MIN_TEXT_CHARS:
        raise CvError("no_text", "Không đọc được chữ trong CV (có thể là ảnh scan).")
    return text


def _read_pdf(data: bytes) -> str:
    import pdfplumber

    try:
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            pages = [(page.extract_text() or "") for page in pdf.pages]
    except Exception as exc:
        raise CvError("corrupt", "Tệp PDF không hợp lệ.") from exc
    return "\n".join(pages)


def _read_docx(data: bytes) -> str:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise CvError("corrupt", "Tệp DOCX không hợp lệ.") from exc
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:  # CV dạng bảng/2 cột
        for row in table.rows:
            for cell in row.cells:
                lines.extend(p.text for p in cell.paragraphs)
    return "\n".join(lines)


# ---------------------------------------------------------------------------
# 2. Chuẩn hoá
# ---------------------------------------------------------------------------

_DECOR = re.compile("[•●▪■□◆◇★☆✓✔➢➤►▶◦‣·●○–—−]+")
_SPACES = re.compile(r"[ \t ]+")


def normalize_text(text: str) -> str:
    """Unicode NFC, bỏ ký tự trang trí, gộp khoảng trắng; giữ nguyên ngắt dòng. Giữ nguyên chữ hoa/thường."""
    text = unicodedata.normalize("NFC", text)
    text = text.replace("–", "-").replace("—", "-")  # dấu gạch dài → gạch ngang trước khi bỏ ký tự trang trí
    text = _DECOR.sub(" ", text)
    lines = [_SPACES.sub(" ", ln).strip() for ln in text.splitlines()]
    return "\n".join(ln for ln in lines if ln)


# ---------------------------------------------------------------------------
# 3. Tách mục (tiêu đề song ngữ)
# ---------------------------------------------------------------------------

SECTION_PATTERNS: dict[str, re.Pattern[str]] = {
    "skills": re.compile(r"^(kỹ năng|kĩ năng|năng lực|chuyên môn|skills?|technical skills?|core competencies)\b"),
    "experience": re.compile(r"^(kinh nghiệm|quá trình làm việc|lịch sử làm việc|work experience|experience|employment)\b"),
    "education": re.compile(r"^(học vấn|trình độ học vấn|giáo dục|education|academic)\b"),
    "projects": re.compile(r"^(dự án|project|projects|personal projects)\b"),
    "summary": re.compile(r"^(mục tiêu|tóm tắt|giới thiệu|objective|summary|profile|about)\b"),
}
_MAX_HEADING_LEN = 40


def _heading_of(line: str) -> Optional[str]:
    cleaned = re.sub(r"[:\-_=*#]+", " ", line.lower()).strip()
    if not cleaned or len(cleaned) > _MAX_HEADING_LEN:
        return None
    for name, pat in SECTION_PATTERNS.items():
        if pat.match(cleaned):
            return name
    return None


def split_sections(text: str) -> dict[str, str]:
    """Tách CV (đã `normalize_text`) thành các mục theo tiêu đề. Phần trước tiêu đề đầu tiên vào `header`."""
    sections: dict[str, list[str]] = {"header": []}
    current = "header"
    for line in text.splitlines():
        heading = _heading_of(line)
        if heading:
            current = heading
            sections.setdefault(current, [])
            # Dòng tiêu đề có nội dung đi kèm: "Kỹ năng: React, Git"
            rest = re.split(r"[:：]", line, maxsplit=1)
            if len(rest) == 2 and rest[1].strip():
                sections[current].append(rest[1].strip())
            continue
        sections[current].append(line)
    return {k: "\n".join(v).strip() for k, v in sections.items() if "\n".join(v).strip()}


# ---------------------------------------------------------------------------
# 4. Số năm kinh nghiệm
# ---------------------------------------------------------------------------

_PRESENT = r"(?:nay|hiện tại|hiện nay|present|now|current)"
_RANGE_MY = re.compile(
    rf"(?<!\d)(\d{{1,2}})\s*[/.\-]\s*(\d{{4}})\s*(?:-|to|đến|tới)\s*(?:(\d{{1,2}})\s*[/.\-]\s*(\d{{4}})|({_PRESENT}))",
    re.IGNORECASE,
)
_RANGE_Y = re.compile(rf"(?<![\d/.\-])(\d{{4}})\s*(?:-|to|đến|tới)\s*(?:(\d{{4}})|({_PRESENT}))(?![\d/])", re.IGNORECASE)
_YEARS_PHRASE = re.compile(
    r"(\d+(?:[.,]\d+)?)\s*\+?\s*(?:năm|years?)\s*(?:kinh nghiệm|of experience|experience)", re.IGNORECASE
)


def _month_index(year: int, month: int) -> int:
    return year * 12 + (month - 1)


def _today_index(today: date) -> int:
    return _month_index(today.year, today.month)


def _intervals(text: str, today: date) -> list[tuple[int, int]]:
    """Các khoảng thời gian [bắt đầu, kết thúc] tính theo chỉ số tháng, kết thúc bao gồm."""
    spans: list[tuple[int, int]] = []
    consumed: list[tuple[int, int]] = []
    now = _today_index(today)
    for m in _RANGE_MY.finditer(text):
        sm, sy = int(m.group(1)), int(m.group(2))
        if not (1 <= sm <= 12 and 1990 <= sy <= today.year + 1):
            continue
        if m.group(5):
            end = now
        else:
            em, ey = int(m.group(3)), int(m.group(4))
            if not (1 <= em <= 12):
                continue
            end = _month_index(ey, em)
        spans.append((_month_index(sy, sm), end))
        consumed.append(m.span())
    for m in _RANGE_Y.finditer(text):
        if any(a <= m.start() < b for a, b in consumed):
            continue
        sy = int(m.group(1))
        if not (1990 <= sy <= today.year + 1):
            continue
        if m.group(3):
            end = now
        else:
            ey = int(m.group(2))
            end = _month_index(ey, 12)
        spans.append((_month_index(sy, 1), end))
    return [(s, e) for s, e in spans if e >= s]


def merge_months(spans: list[tuple[int, int]]) -> int:
    """Tổng số tháng sau khi gộp các khoảng chồng nhau (không tính trùng)."""
    total = 0
    cur_s: Optional[int] = None
    cur_e = 0
    for s, e in sorted(spans):
        if cur_s is None:
            cur_s, cur_e = s, e
        elif s <= cur_e + 1:
            cur_e = max(cur_e, e)
        else:
            total += cur_e - cur_s + 1
            cur_s, cur_e = s, e
    if cur_s is not None:
        total += cur_e - cur_s + 1
    return total


def extract_years_exp(text: str, sections: Optional[dict[str, str]] = None, today: Optional[date] = None) -> float:
    """Số năm kinh nghiệm: cộng dồn các khoảng thời gian không chồng nhau trong mục kinh nghiệm;
    nếu không có thì dùng cụm "X năm kinh nghiệm"."""
    today = today or date.today()
    sections = sections if sections is not None else split_sections(text)
    if "experience" in sections:
        scope = sections["experience"]
    else:  # không có mục kinh nghiệm: bỏ phần học vấn để không đếm nhầm thời gian học
        scope = "\n".join(v for k, v in sections.items() if k != "education")
    months = merge_months(_intervals(scope, today))
    if months > 0:
        return round(months / 12, 1)
    m = _YEARS_PHRASE.search(text)
    if m:
        return round(float(m.group(1).replace(",", ".")), 1)
    return 0.0


# ---------------------------------------------------------------------------
# 5. Học vấn (thang 1-4)
# ---------------------------------------------------------------------------

_EDU_LEVELS: list[tuple[int, re.Pattern[str]]] = [
    (4, re.compile(r"thạc sĩ|thạc sỹ|\bmaster|\bm\.?sc\b|\bmba\b|tiến sĩ|\bph\.?d\b", re.IGNORECASE)),
    (3, re.compile(r"đại học|cử nhân|kỹ sư|\buniversity\b|\bbachelor|\bb\.?sc\b|học viện", re.IGNORECASE)),
    (2, re.compile(r"cao đẳng|\bcollege\b|\bassociate\b", re.IGNORECASE)),
    (1, re.compile(r"trung cấp|\bvocational\b", re.IGNORECASE)),
]


def extract_education_level(text: str, sections: Optional[dict[str, str]] = None) -> int:
    """Trình độ cao nhất tìm thấy, thang 1-4 (1 Trung cấp, 2 Cao đẳng, 3 Đại học, 4 Thạc sĩ+); 0 nếu không rõ."""
    sections = sections if sections is not None else split_sections(text)
    scope = sections.get("education") or text
    for level, pat in _EDU_LEVELS:
        if pat.search(scope):
            return level
    return 0


# ---------------------------------------------------------------------------
# 6. Ghép các bước
# ---------------------------------------------------------------------------


def parse_cv(data: bytes, filename: str, index=None, embedder=None, today: Optional[date] = None) -> dict:
    """Đọc CV và trả về kết quả phân tích (lưu vào `resumes.parsed_json`).

    Trả về: {text, skills:[{name,source,confidence}], years_exp, education_level, sections, embedding(bytes)}.
    """
    from .embedder import get_embedder, to_bytes
    from .extract import extract_skills

    raw = extract_text(data, filename)
    text = normalize_text(raw)
    sections = split_sections(text)
    hits = extract_skills(text, index=index, embedder=embedder, skills_section=sections.get("skills", ""))
    emb = (embedder or get_embedder()).encode([text[:3000]])[0]
    return {
        "text": text,
        "skills": [h.to_dict() for h in hits],
        "years_exp": extract_years_exp(text, sections, today),
        "education_level": extract_education_level(text, sections),
        "sections": sections,
        "embedding": to_bytes(emb),
    }


def suggest_skills(text: str, index=None) -> list[dict]:
    """Gợi ý kỹ năng từ mô tả công việc (JD) cho nhà tuyển dụng; dùng tầng 1-2."""
    from .extract import extract_skills

    normalized = normalize_text(text)
    return [h.to_dict() for h in extract_skills(normalized, index=index, tiers=(1, 2))]
